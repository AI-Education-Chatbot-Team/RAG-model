"""
Integration tests that run the REAL MarkItDown parser and embedding model on real
files. They download the embedding model the first time (~90 MB) and are slower, so they
are skipped by default. Run them with:

    pytest -m integration

Supabase and Groq are still faked; no API keys are needed.
"""
import pytest

pytestmark = pytest.mark.integration


@pytest.fixture(scope="module")
def real_ingest():
    """Swap the fakes from conftest for the real converter and embedding model."""
    import ingest
    import shared
    from conftest import REAL

    old = (ingest.md, shared.embed_model)
    ingest.md = REAL["markitdown.MarkItDown"]()
    shared.embed_model = REAL["sentence_transformers.SentenceTransformer"]("all-MiniLM-L6-v2")
    yield ingest
    ingest.md, shared.embed_model = old


@pytest.fixture(scope="module")
def samples(tmp_path_factory):
    d = tmp_path_factory.mktemp("samples")

    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(d / "two_pages.pdf"))
    c.drawString(72, 720, "Luke Skywalker is a Jedi from Tatooine.")
    c.showPage()
    c.drawString(72, 720, "Page two talks about Darth Vader.")
    c.save()

    import docx
    doc = docx.Document()
    doc.add_heading("Course Notes", 1)
    doc.add_paragraph("Photosynthesis converts light into chemical energy.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Term", "Meaning"
    table.cell(1, 0).text, table.cell(1, 1).text = "ATP", "Energy molecule"
    doc.save(d / "notes.docx")

    (d / "notes.txt").write_text("The mitochondria is the powerhouse of the cell.\n")
    (d / "notes.md").write_text("# Heading\n\nMarkdown body text about binning.\n")
    (d / "grades.csv").write_text("name,grade\nAlice,A\nBob,B\n")
    (d / "empty.txt").write_text("")
    (d / "unicode.txt").write_text("Café naïve résumé — 日本語 テキスト 😀\n")
    (d / "broken.pdf").write_bytes(b"%PDF-1.4 this is not really a pdf")
    (d / "big.txt").write_text("".join(f"Sentence number {i} about data mining. " for i in range(3000)))
    return d


@pytest.mark.parametrize("name, expected", [
    ("two_pages.pdf", "Luke Skywalker is a Jedi"),
    ("notes.docx", "Photosynthesis converts light"),
    ("notes.txt", "powerhouse of the cell"),
    ("notes.md", "Markdown body text"),
    ("grades.csv", "Alice"),
])
def test_each_supported_file_type_is_extracted(real_ingest, samples, name, expected):
    chunks = real_ingest.extract_and_chunk(str(samples / name), name)

    assert chunks, f"no chunks extracted from {name}"
    assert expected in " ".join(c["text"] for c in chunks)
    assert all(c["source_file"] == name for c in chunks)


def test_docx_table_contents_are_kept(real_ingest, samples):
    text = real_ingest.extract_and_chunk(str(samples / "notes.docx"), "notes.docx")[0]["text"]
    assert "ATP" in text and "Energy molecule" in text


def test_unicode_survives_extraction(real_ingest, samples):
    text = real_ingest.extract_and_chunk(str(samples / "unicode.txt"), "unicode.txt")[0]["text"]
    assert "Café" in text and "日本語" in text


def test_empty_file_gives_no_chunks(real_ingest, samples):
    assert real_ingest.extract_and_chunk(str(samples / "empty.txt"), "empty.txt") == []


def test_large_file_is_split_into_bounded_chunks(real_ingest, samples):
    chunks = real_ingest.extract_and_chunk(str(samples / "big.txt"), "big.txt")

    assert len(chunks) > 50
    assert all(len(c["text"]) <= 1000 for c in chunks)
    joined = " ".join(c["text"] for c in chunks)
    assert "Sentence number 0 " in joined and "Sentence number 2999 " in joined


@pytest.mark.xfail(strict=True, raises=AssertionError, reason="BUG: MarkItDown falls back to plain text, so a broken PDF is stored as raw PDF bytes")
def test_broken_pdf_is_not_stored_as_raw_text(real_ingest, samples):
    try:
        chunks = real_ingest.extract_and_chunk(str(samples / "broken.pdf"), "broken.pdf")
    except Exception:
        return
    assert not any("%PDF" in c["text"] for c in chunks)


def test_real_embeddings_match_database_dimension(real_ingest):
    import shared
    vector = shared.embed_texts(["hello world"])[0]
    assert len(vector) == 384


def test_related_text_scores_higher_than_unrelated(real_ingest):
    import numpy as np
    import shared
    q, near, far = shared.embed_texts([
        "What does the mitochondria do?",
        "The mitochondria is the powerhouse of the cell.",
        "Luke Skywalker is a Jedi from Tatooine.",
    ])
    cos = lambda a, b: float(np.dot(a, b) / (np.linalg.norm(a) * np.linalg.norm(b)))
    assert cos(q, near) > cos(q, far)